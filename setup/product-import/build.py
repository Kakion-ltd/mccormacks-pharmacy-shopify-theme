"""Build the product import template for adding products by CSV.

    python3 setup/product-import/build.py

Reads the live store (read-only, through the logged-in Shopify CLI) and writes:

  product-import-template.csv   the file Shopify imports: headers and one example row
  product-import-workbook.xlsx  the same columns to fill in, with pick-lists, plus the
                                valid types, category tags, vendors and existing products

The lists are a snapshot. Types and tags come from the live collection rules, so when a
category page is added, removed or re-ruled, rebuild before sending the workbook out
again. A copy that is not rebuilt is how a product ends up in no category.
"""
import csv
import datetime
import json
import os
import subprocess
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
STORE = os.environ.get("STORE", "mccormackpharmacy.myshopify.com")
TODAY = datetime.date.today().strftime("%-d %B %Y")
ROWS = 500  # rows given pick-lists and text formatting on the Products sheet

# Shopify's current product CSV headers, checked against
# help.shopify.com/en/manual/products/import-export/using-csv on 25 Sep 2026. Only the
# columns this store uses: every product is single-variant, so no option columns.
COLUMNS = [
    "Title", "URL handle", "Description", "Vendor", "Type", "Tags",
    "Published on online store", "Status", "SKU", "Barcodes", "Price",
    "Compare-at price", "Charge tax", "Inventory tracker", "Inventory quantity",
    "Continue selling when out of stock", "Requires shipping", "Product image URL",
    "Image alt text",
]
# Defaults match every existing product: taxed, stock tracked, no sales when out of stock.
EXAMPLE = {
    "Title": "EXAMPLE - DELETE THIS ROW - Aveeno Daily Moisturising Lotion 200ml",
    "URL handle": "",
    "Description": "<p>A daily moisturiser for dry skin, with colloidal oatmeal.</p>",
    "Vendor": "Aveeno",
    "Type": "Skincare > Body Care",
    "Tags": "Dry Skin, Moisturisers",
    "Published on online store": "TRUE",
    "Status": "active",
    "SKU": "5012345678900",
    "Barcodes": "",
    "Price": "12.99",
    "Compare-at price": "",
    "Charge tax": "TRUE",
    "Inventory tracker": "shopify",
    "Inventory quantity": "12",
    "Continue selling when out of stock": "FALSE",
    "Requires shipping": "TRUE",
    "Product image URL": "",
    "Image alt text": "",
}
NOTES = {
    "Title": "Required. The product name as shoppers see it.",
    "URL handle": "Leave blank: Shopify makes it from the title.",
    "Description": "Plain text or simple HTML (<p>, <ul><li>).",
    "Vendor": "Pick from the list. It must match an existing vendor exactly, or the product "
              "misses its brand page and appears under a new, separate brand.",
    "Type": "Pick from the list. Anything else and the product lands on no category page.",
    "Tags": "Comma-separated. Copy tags exactly from the 'Category tags' sheet.",
    "SKU": "The product's barcode (EAN, usually 13 digits). The whole store keeps the "
           "barcode here. This column is formatted as text so Excel keeps every digit.",
    "Barcodes": "Leave blank. The barcode goes in SKU.",
    "Compare-at price": "Only for a reduced product: the old, higher price. This is what "
                        "shows the SALE badge. Blank for a normal product.",
    "Inventory quantity": "Stock count. See the Read me: this may be ignored if the store "
                          "has more than one location.",
}


def query(q):
    env = dict(os.environ, CI="1", SHOPIFY_CLI_NO_ANALYTICS="1")
    out = subprocess.run(["npx", "shopify", "store", "execute", "-s", STORE, "-q", q],
                         cwd=ROOT, env=env, capture_output=True, text=True, check=True).stdout
    return json.loads(out[out.index("{"):] if out.lstrip().startswith("{") else out[out.index("\n{") + 1:])


def paged(template, key):
    items, after = [], None
    while True:
        node = query(template.replace("AFTER", f', after: "{after}"' if after else ""))[key]
        items += [e["node"] for e in node["edges"]]
        if not node["pageInfo"]["hasNextPage"]:
            return items
        after = node["pageInfo"]["endCursor"]


def fetch():
    products = paged('{ products(first: 250AFTER) { edges { node { title vendor productType '
                     'variants(first: 1) { edges { node { sku } } } } } '
                     'pageInfo { hasNextPage endCursor } } }', "products")
    collections = paged('{ collections(first: 250AFTER) { edges { node { handle title '
                        'productsCount { count } ruleSet { rules { column relation condition } } } } '
                        'pageInfo { hasNextPage endCursor } } }', "collections")
    return products, collections


def lists(products, collections):
    taxonomy = json.load(open(os.path.join(ROOT, "setup", "taxonomy.json")))
    by_title = {c["title"].lower(): c for c in collections}
    rules = lambda c: (c["ruleSet"] or {}).get("rules") or []

    # Department pages match a type prefix; group pages match an exact type.
    prefixes = {r["condition"]: c["title"] for c in collections for r in rules(c)
                if r["column"] == "TYPE" and r["relation"] == "STARTS_WITH"}
    pages_for = defaultdict(list)
    for c in collections:
        for r in rules(c):
            if r["column"] == "TYPE" and r["relation"] == "EQUALS":
                pages_for[r["condition"]].append(c["title"])
    types = []
    for t in sorted(pages_for):
        dept = [d for p, d in prefixes.items() if t.startswith(p + " >") or t == p]
        if dept:  # a type with no department prefix skips its department page
            types.append((t, dept[0], ", ".join(sorted(set(pages_for[t])))))
    # Bare department names: a fallback that reaches the department page only.
    for p in ["Pharmacy", "Supplements", "Skincare", "Beauty", "Toiletries", "Baby", "Gifts"]:
        if p in prefixes:
            types.append((p, prefixes[p], "(department page only)"))

    # Sub-category pages match a tag. Take the tag from the live rule, not the page title.
    tags = []
    for m in taxonomy:
        entries = [(g["title"], g["title"]) for g in m["groups"]] + \
                  [(g["title"], i) for g in m["groups"] for i in g["items"]] + \
                  [("", i) for i in m["flat"]]
        for group, title in entries:
            c = by_title.get(title.lower())
            for r in rules(c) if c else []:
                if r["column"] == "TAG":
                    tags.append((m["menu"], group if group != title else "", title,
                                 r["condition"], c["productsCount"]["count"]))
    special = [(c["title"], r["condition"], c["productsCount"]["count"]) for c in collections
               for r in rules(c) if r["column"] == "TAG" and c["handle"] in
               ("sale", "hot-offers", "new-in", "pharmacist-review-required")]

    brand_pages = {r["condition"] for c in collections for r in rules(c) if r["column"] == "VENDOR"}
    vendors = sorted(Counter(p["vendor"] for p in products if p["vendor"]).items(),
                     key=lambda v: v[0].lower())
    existing = sorted(((p["title"], (p["variants"]["edges"] or [{"node": {"sku": ""}}])[0]["node"]["sku"] or "",
                        p["vendor"], p["productType"]) for p in products), key=lambda r: r[0].lower())
    return types, tags, special, brand_pages, vendors, existing


def write_csv():
    with open(os.path.join(HERE, "product-import-template.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(COLUMNS)
        w.writerow([EXAMPLE[c] for c in COLUMNS])


BOLD, HEAD_FILL = Font(name="Arial", bold=True), PatternFill("solid", fgColor="D9E7C6")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
ARIAL = Font(name="Arial")


def table(ws, headers, rows, widths):
    ws.append(headers)
    for cell in ws[1]:
        cell.font, cell.fill = BOLD, HEAD_FILL
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = ARIAL
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_workbook(types, tags, special, brand_pages, vendors, existing):
    wb = Workbook()

    readme = wb.active
    readme.title = "Read me"
    readme.column_dimensions["A"].width = 110
    lines = [
        ("Adding products to McCormack's by CSV", True),
        (f"Lists in this workbook were taken from the live store on {TODAY}.", False),
        ("", False),
        ("How to fill it in", True),
        ("1. Fill in the Products sheet, one row per product. Yellow cells are yours to edit.", False),
        ("2. Delete the example row (row 2) before you save.", False),
        ("3. Save the Products sheet as 'CSV UTF-8 (Comma delimited)'. Excel saves only the "
         "sheet you are on, so be on Products when you save.", False),
        ("4. Import in Shopify: Products > Import. Leave 'Overwrite products with matching "
         "handles' UNticked.", False),
        ("", False),
        ("Rules that will catch you out", True),
        ("TYPE must be picked from the list, spelled exactly. A type that is not on the list "
         "puts the product on no category page. Shoppers can then only find it by searching.", False),
        ("TAGS put a product on the smaller category pages (for example Cough, Sore Throat, "
         "Acne & Blemish). Copy them from the 'Category tags' sheet, spelled exactly, "
         "separated by commas. The type alone does not reach these pages.", False),
        ("VENDOR must match the store's existing spelling exactly (pick from the list), even "
         "where the store's spelling is not the brand's own. The store has 'A.VOGEL'; typing "
         "'A. Vogel' makes a second, separate brand, and a product under it is missing from "
         "the brand's page and filters. A genuinely new brand is allowed: check the Vendors "
         "sheet first, then use one spelling for every product of that brand.", False),
        ("BARCODE goes in the SKU column, not the Barcodes column. That is how every product "
         "already on the store is set up. The SKU column is formatted as text so Excel does "
         "not turn 5012345678900 into 5.01E+12 or drop a leading zero.", False),
        ("SALE: the SALE badge and the crossed-out price appear only when 'Compare-at price' "
         "is filled in AND higher than 'Price'. Put the old price there. Leave it blank for "
         "anything not reduced. To also list the product on the Sale page, add the tag 'sale'.", False),
        ("MEDICINES: any product the pharmacist must review needs the tag 'pharmacist-review'. "
         "Without it the product can be added to the bag in one click from any listing. The "
         "pharmacist questionnaire itself is set on the product in Shopify admin after import; "
         "it cannot be set from this file. Ask the pharmacist which products need either.", False),
        ("CHECK IT IS NEW: search the 'On the store already' sheet for the barcode before "
         "adding a row. A product imported twice appears twice on the site.", False),
        ("STOCK: 'Inventory quantity' only works if the store keeps stock in a single "
         "location. The test import below will show whether it took.", False),
        ("", False),
        ("Import in two steps", True),
        ("1. Test: import a file with just three products, chosen to cover different types.", False),
        ("2. Check each of the three on the website: it opens; the price is right; it is on "
         "its category page (from Type) and its smaller category page (from Tags); it is on "
         "its brand page; stock shows as in stock; a reduced product shows SALE; a medicine "
         "has no quick Add button on listings.", False),
        ("3. Fix anything wrong in the spreadsheet, not only in Shopify, so the full file "
         "carries the fix.", False),
        ("4. Then import the rest, in batches of about 100, checking a few from each batch.", False),
    ]
    for text, bold in lines:
        readme.append([text])
        cell = readme.cell(row=readme.max_row, column=1)
        cell.font = Font(name="Arial", bold=bold, size=13 if bold else 11)
        cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws = wb.create_sheet("Products")
    ws.append(COLUMNS)
    ws.append([EXAMPLE[c] for c in COLUMNS])
    for i, c in enumerate(COLUMNS, start=1):
        head = ws.cell(row=1, column=i)
        head.font, head.fill = BOLD, HEAD_FILL
        if c in NOTES:
            head.comment = Comment(NOTES[c], "McCormack's")
        # Styled by column, not by cell: Excel exports every cell that carries formatting,
        # so 500 pre-formatted rows would save as 500 blank CSV lines for Shopify to reject.
        # A column style reaches the cells Keelan types into without making them "used".
        dim = ws.column_dimensions[head.column_letter]
        dim.width = max(14, min(40, len(c) + 4))
        dim.font, dim.fill = ARIAL, INPUT_FILL
        if c in ("SKU", "Barcodes"):
            dim.number_format = "@"
        cell = ws.cell(row=2, column=i)
        cell.font, cell.fill = ARIAL, INPUT_FILL
        if c in ("SKU", "Barcodes"):
            cell.number_format = "@"
    ws.freeze_panes = "B2"
    col = {c: ws.cell(row=1, column=i + 1).column_letter for i, c in enumerate(COLUMNS)}
    rng = lambda c: f"{col[c]}2:{col[c]}{ROWS + 1}"

    def pick(source, c, strict, message):
        dv = DataValidation(type="list", formula1=source, allow_blank=True,
                            errorStyle="stop" if strict else "warning",
                            errorTitle=c, error=message)
        ws.add_data_validation(dv)
        dv.add(rng(c))

    pick(f"='Types'!$A$2:$A${len(types) + 1}", "Type", True,
         "Pick a type from the list. Any other type puts the product on no category page.")
    pick(f"='Vendors'!$A$2:$A${len(vendors) + 1}", "Vendor", False,
         "This vendor is not on the store yet. Check the spelling against the Vendors "
         "sheet. Continue only if this is a genuinely new brand.")
    for c, opts in (("Published on online store", "TRUE,FALSE"), ("Status", "active,draft"),
                    ("Charge tax", "TRUE,FALSE"), ("Continue selling when out of stock", "FALSE,TRUE"),
                    ("Requires shipping", "TRUE,FALSE"), ("Inventory tracker", "shopify")):
        pick(f'"{opts}"', c, True, f"Pick {opts.replace(',', ' or ')}.")

    table(wb.create_sheet("Types"), ["Type (pick exactly)", "Department page", "Also appears on"],
          types, [44, 24, 60])
    table(wb.create_sheet("Category tags"),
          ["Menu", "Group", "Category page", "Tag to add (exact)", "Products there now"],
          tags, [22, 30, 40, 40, 20])
    t = wb["Category tags"]
    t.append([])
    t.append(["Special tags", "", "Page / effect", "Tag to add (exact)", "Products there now"])
    for cell in t[t.max_row]:
        cell.font, cell.fill = BOLD, HEAD_FILL
    for title, tag, n in special:
        t.append(["", "", title, tag, n])
    table(wb.create_sheet("Vendors"), ["Vendor (exact spelling)", "Products on store", "Has a brand page"],
          [(v, n, "yes" if v in brand_pages else "") for v, n in vendors], [40, 18, 18])
    table(wb.create_sheet("On the store already"), ["Title", "SKU (barcode)", "Vendor", "Type"],
          existing, [60, 18, 28, 36])
    for row in wb["On the store already"].iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "@"
    wb.active = 1
    wb.save(os.path.join(HERE, "product-import-workbook.xlsx"))


if __name__ == "__main__":
    data = lists(*fetch())
    write_csv()
    write_workbook(*data)
    types, tags, special, brand_pages, vendors, existing = data
    print(f"{len(types)} types, {len(tags)} category tags, {len(special)} special tags, "
          f"{len(vendors)} vendors ({len(brand_pages)} with brand pages), {len(existing)} existing products")
